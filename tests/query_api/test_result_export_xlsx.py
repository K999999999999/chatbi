from __future__ import annotations

from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from src.query_api.export_xlsx import _verify_workbook, create_query_workbook


def test_query_workbook_verifier_rejects_dtd_and_external_entities(tmp_path):
    destination = tmp_path / "unsafe.xlsx"
    with ZipFile(destination, "w") as archive:
        archive.writestr(
            "xl/workbook.xml",
            b'<?xml version="1.0"?><!DOCTYPE workbook '
            b'[<!ENTITY secret SYSTEM "file:///etc/passwd">]>'
            b"<workbook>&secret;</workbook>",
        )

    with pytest.raises(ValueError, match="invalid or unsafe XML"):
        _verify_workbook(destination, data_rows=1, data_columns=1, note_rows=1)


def test_query_workbook_preserves_raw_values_and_cell_types(tmp_path):
    destination = tmp_path / "result.xlsx"
    create_query_workbook(
        {
            "question": "收入 = 真实文本",
            "result_time": None,
            "saved_time": "2026-10-05T11:00:00+00:00",
            "export_time": "2026-10-06T11:00:00+00:00",
            "result": {
                "columns": ["指标", "编号", "说明", "链接"],
                "rows": [
                    [
                        "12345678901234567890.1234",
                        "00123",
                        None,
                        "https://example.com/report",
                    ],
                    ["=1+1", 0, "", ""],
                    ["〈NULL〉", "〈空字符串〉", "原始文本", "javascript:alert(1)"],
                ],
                "row_count": 3,
                "truncated": False,
                "result_metadata": None,
            },
        },
        destination,
    )

    workbook = load_workbook(destination, data_only=False)
    assert workbook.sheetnames == ["原始数据", "结果说明"]
    data = workbook["原始数据"]
    assert list(data.values) == [
        ("指标", "编号", "说明", "链接"),
        (
            "12345678901234567890.1234",
            "00123",
            "〈NULL〉",
            "https://example.com/report",
        ),
        ("=1+1", 0, "〈空字符串〉", "〈空字符串〉"),
        ("〈NULL〉", "〈空字符串〉", "原始文本", "javascript:alert(1)"),
    ]
    assert data["A3"].data_type == "s"
    assert data["B2"].data_type == "s"
    assert data["B3"].data_type == "n"
    assert data["D2"].hyperlink is None
    assert data["D4"].hyperlink is None
    notes = list(workbook["结果说明"].values)
    assert ("来源问题", "收入 = 真实文本") in notes
    assert ("结果时间", "原结果完成时间未保存") in notes
    assert ("成果保存时间", "2026-10-05T11:00:00+00:00") in notes
    assert ("C2 · 说明", "null · null") in notes
    assert ("C3 · 说明", 'string · ""') in notes
    assert ("A4 · 指标", 'string · "〈NULL〉"') in notes
    assert ("B4 · 编号", 'string · "〈空字符串〉"') in notes
    assert ("D3 · 链接", 'string · ""') in notes
    assert not any(cell.data_type == "f" for row in data.iter_rows() for cell in row)
    workbook.close()


def test_query_workbook_rejects_unrepresentable_cell_text(tmp_path):
    import pytest

    with pytest.raises(ValueError, match="单元格长度"):
        create_query_workbook(
            {
                "question": "超长",
                "result_time": None,
                "export_time": "2026-10-06T11:00:00+00:00",
                "result": {
                    "columns": ["说明"],
                    "rows": [["x" * 32768]],
                    "row_count": 1,
                    "truncated": False,
                    "result_metadata": None,
                },
            },
            tmp_path / "too-long.xlsx",
        )


def test_query_workbook_applies_excel_utf16_and_numeric_precision_limits(tmp_path):
    import pytest

    destination = tmp_path / "limits.xlsx"
    create_query_workbook(
        {
            "question": "边界值",
            "result_time": None,
            "export_time": "2026-10-06T11:00:00+00:00",
            "result": {
                "columns": [
                    "安全数值",
                    "高精度数值",
                    "Excel 数值下溢值",
                    "Excel 最小数值",
                    "Excel 超大范围",
                ],
                "rows": [
                    [
                        123456789012345.0,
                        1234567890123456.0,
                        1e-320,
                        2.2251e-308,
                        1e308,
                    ]
                ],
                "row_count": 1,
                "truncated": False,
                "result_metadata": None,
            },
        },
        destination,
    )
    workbook = load_workbook(destination, read_only=True)
    row = list(workbook["原始数据"].values)[1]
    assert row == (
        123456789012345.0,
        "1234567890123456.0",
        "1e-320",
        2.2251e-308,
        "1e+308",
    )
    notes = list(workbook["结果说明"].values)
    assert ("B2 · 高精度数值", "number-text · 1234567890123456.0") in notes
    assert ("C2 · Excel 数值下溢值", "number-text · 1e-320") in notes
    assert ("E2 · Excel 超大范围", "number-text · 1e+308") in notes
    workbook.close()

    with pytest.raises(ValueError, match="单元格长度"):
        create_query_workbook(
            {
                "question": "超长 UTF-16 内容",
                "result_time": None,
                "export_time": "2026-10-06T11:00:00+00:00",
                "result": {
                    "columns": ["说明"],
                    "rows": [["😀" * 16384]],
                    "row_count": 1,
                    "truncated": False,
                    "result_metadata": None,
                },
            },
            tmp_path / "too-long-utf16.xlsx",
        )


def test_query_workbook_preserves_empty_and_truncated_snapshot_semantics(tmp_path):
    empty_path = tmp_path / "empty.xlsx"
    create_query_workbook(
        {
            "question": "没有匹配结果",
            "result_time": None,
            "export_time": "2026-10-06T11:00:00+00:00",
            "result": {
                "columns": ["月份", "销售额"],
                "rows": [],
                "row_count": 0,
                "truncated": False,
                "result_metadata": None,
            },
        },
        empty_path,
    )
    empty = load_workbook(empty_path, read_only=True)
    assert list(empty["原始数据"].values) == [("月份", "销售额")]
    assert ("结果行数", "0") in list(empty["结果说明"].values)
    empty.close()

    truncated_path = tmp_path / "truncated.xlsx"
    create_query_workbook(
        {
            "question": "返回结果已截断",
            "result_time": None,
            "export_time": "2026-10-06T11:00:00+00:00",
            "result": {
                "columns": ["月份"],
                "rows": [["2026-02"], ["2026-01"]],
                "row_count": 10,
                "truncated": True,
                "result_metadata": None,
            },
        },
        truncated_path,
    )
    truncated = load_workbook(truncated_path, read_only=True)
    assert list(truncated["原始数据"].values) == [
        ("月份",),
        ("2026-02",),
        ("2026-01",),
    ]
    assert (
        "截断提示",
        "仅包含快照当前返回行，未重新查询或获取全量数据。",
    ) in list(truncated["结果说明"].values)
    truncated.close()


def test_query_workbook_keeps_an_empty_source_column_name_empty(tmp_path):
    destination = tmp_path / "empty-column.xlsx"
    create_query_workbook(
        {
            "question": "空列名",
            "result_time": None,
            "export_time": "2026-10-06T11:00:00+00:00",
            "result": {
                "columns": ["", "正常列名"],
                "rows": [[1, 2]],
                "row_count": 1,
                "truncated": False,
                "result_metadata": None,
            },
        },
        destination,
    )
    workbook = load_workbook(destination, read_only=True)
    assert list(workbook["原始数据"].values) == [("", "正常列名"), (1, 2)]
    workbook.close()
