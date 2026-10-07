"""容器验收的专用账号与证据；只在一次性工具中执行。"""

import argparse
import hashlib
import json
import os
import pty
import secrets
import select as io_select
import struct
import subprocess
import sys
import time
import unicodedata
import zlib
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from sqlalchemy import func, select, text
from sqlalchemy.orm import Session, sessionmaker

from src.authorization import hash_password
from src.authorization.auth_service import AuthService
from src.chatbi_control.database import ControlDatabaseConfig, create_control_engine
from src.chatbi_control.models import Role, User, UserSession
from tests.browser_real_reference import reference

REPORT = Path("/reports")


def prepare():
    engine = create_control_engine(
        ControlDatabaseConfig.from_environment(require_migrator=False)
    )
    username = f"web-e2e-container-{uuid4().hex}"
    password = secrets.token_hex(24)
    try:
        with Session(engine) as session, session.begin():
            role = session.scalar(select(Role).where(Role.name == "analyst"))
            if role is None:
                raise RuntimeError("缺少 analyst 角色，先显式 migrate")
            user = User(
                username=username,
                password_hash=hash_password(password),
                is_active=True,
                must_change_password=False,
                roles=[role],
            )
            session.add(user)
            session.flush()
            user_id = user.id
        REPORT.mkdir(exist_ok=True)
        account = REPORT / "account.json"
        account.write_text(json.dumps({"username": username, "id": user_id}))
        reference_data = reference()
        credentials = REPORT / "credentials.env"
        credentials.touch(mode=0o600)
        credentials.write_text(
            f"CHATBI_REAL_E2E_USERNAME={username}\nCHATBI_REAL_E2E_PASSWORD={password}\n"
            f"CHATBI_CONTAINER_REFERENCE={json.dumps(reference_data, ensure_ascii=False)}\n"
        )
    finally:
        engine.dispose()
    print("容器验收账号与只读参考已准备。")


def cleanup():
    account = json.loads((REPORT / "account.json").read_text())
    if not account["username"].startswith("web-e2e-container-"):
        raise RuntimeError("禁止修改非本次验收账号")
    engine = create_control_engine(
        ControlDatabaseConfig.from_environment(require_migrator=False)
    )
    try:
        with Session(engine) as session:
            user = session.get(User, account["id"])
            if user is None or user.username != account["username"]:
                raise RuntimeError("账号身份不一致")
        AuthService(sessionmaker(engine)).disable_user(account["id"])
        with Session(engine) as session:
            disabled = not session.get(User, account["id"]).is_active
            active = session.scalar(
                select(func.count())
                .select_from(UserSession)
                .where(
                    UserSession.user_id == account["id"],
                    UserSession.revoked_at.is_(None),
                )
            )
        result = {"disabled": disabled, "active_sessions": active}
        (REPORT / "cleanup.json").write_text(json.dumps(result))
        if not disabled or active != 0:
            raise RuntimeError("验收账号未完整清理")
        browser = REPORT / "browser.json"
        if browser.exists():
            evidence = json.loads(browser.read_text())
            evidence["cleanup"] = result
            metadata = REPORT / "runtime.json"
            if metadata.exists():
                evidence["runtime"] = json.loads(metadata.read_text())
            hot_reload = REPORT / "hot-reload.json"
            if hot_reload.exists():
                evidence["hot_reload"] = json.loads(hot_reload.read_text())
            recovery = REPORT / "recovery.json"
            if recovery.exists():
                evidence["execution_recovery"] = json.loads(recovery.read_text())
            exports = REPORT / "export-verification.json"
            if exports.exists():
                evidence["export_verification"] = json.loads(exports.read_text())
            images = REPORT / "images.jsonl"
            if images.exists():
                evidence["image_ids"] = [
                    json.loads(line) for line in images.read_text().splitlines()
                ]
            evidence["suite_status"] = evidence.get("status", "failed")
            browser.write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
    finally:
        engine.dispose()
        (REPORT / "credentials.env").unlink(missing_ok=True)
    print("容器验收账号已禁用，活跃 Session 为0。")


def initialize_admin():
    if os.environ.get("CHATBI_CONTAINER_ISOLATED") != "1":
        raise RuntimeError("首次管理员验收仅允许隔离环境")
    master, slave = pty.openpty()
    password = secrets.token_hex(24)
    process = subprocess.Popen(
        [
            sys.executable,
            "-m",
            "src.bootstrap",
            "create-admin",
            "--username",
            "container-check-admin",
        ],
        stdin=slave,
        stdout=slave,
        stderr=slave,
    )
    os.close(slave)
    captured = b""
    prompts = ["首个管理员密码", "再次输入首个管理员密码"]
    deadline = time.monotonic() + 30
    try:
        while process.poll() is None and time.monotonic() < deadline:
            if io_select.select([master], [], [], 0.2)[0]:
                try:
                    captured += os.read(master, 4096)
                except OSError:
                    break
                if prompts and prompts[0].encode() in captured:
                    os.write(master, (password + "\n").encode())
                    prompts.pop(0)
                    captured = b""
        if process.wait(timeout=5) != 0 or prompts:
            raise RuntimeError("交互管理员创建未成功；诊断不回显密码或终端记录")
    finally:
        if process.poll() is None:
            process.kill()
            process.wait()
        os.close(master)
    print("隔离环境交互管理员创建通过。")


def expire_analysis():
    account = json.loads((REPORT / "account.json").read_text())
    evidence = json.loads((REPORT / "browser.json").read_text())
    engine = create_control_engine(
        ControlDatabaseConfig.from_environment(require_migrator=False)
    )
    try:
        with engine.begin() as connection:
            changed = connection.execute(
                text(
                    "UPDATE business_analysis_runs SET expires_at=CURRENT_TIMESTAMP-INTERVAL '1 second' WHERE analysis_run_id=:run AND owner_subject=:owner"
                ),
                {
                    "run": evidence["restart_inputs"]["analysis_run_id"],
                    "owner": "local:" + account["username"],
                },
            )
            # 身份subject沿用LocalSessionIdentityProvider；先核实唯一归属再推进时间。
            if changed.rowcount != 1:
                raise RuntimeError("验收分析身份不一致，未完成TTL检查")
    finally:
        engine.dispose()
    print("专用验收分析检查点已过期，完成历史快照保留。")


def verify_persistence():
    evidence = json.loads((REPORT / "browser.json").read_text())
    current = reference()
    for key in ("net_sales", "runtime_identity"):
        if current[key] != evidence["reference"][key]:
            raise RuntimeError("停止重启后的业务结果或RAG资产身份改变")
    print("停止重启后的业务数据与RAG资产身份保持。")


def acceptance_snapshot():
    """Record only counts and Seed identity around an isolated API restart."""
    import psycopg

    control = create_control_engine(
        ControlDatabaseConfig.from_environment(require_migrator=False)
    )
    try:
        with control.connect() as connection:
            state = {
                name: connection.execute(
                    text(f"SELECT count(*) FROM {table}")
                ).scalar_one()
                for name, table in (
                    ("users", "users"),
                    ("histories", "history_records"),
                    ("turns", "history_turns"),
                    ("saved_results", "saved_results"),
                )
            }
    finally:
        control.dispose()

    with psycopg.connect(
        host=os.environ["POSTGRES_HOST"],
        port=int(os.environ.get("POSTGRES_PORT", "5432")),
        dbname=os.environ["POSTGRES_DB"],
        user=os.environ["POSTGRES_APP_USER"],
        password=os.environ["POSTGRES_APP_PASSWORD"],
        connect_timeout=5,
    ) as connection:
        connection.read_only = True
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT seed_version, (SELECT count(*) FROM mart_sales.fct_sales_order_line) "
                "FROM mart_sales.dev_seed_metadata WHERE singleton"
            )
            row = cursor.fetchone()
    if row is None:
        raise RuntimeError("隔离验收 Sales Mart Seed 身份缺失")
    state["business_seed_version"] = row[0]
    state["business_fact_rows"] = row[1]
    print(json.dumps(state, sort_keys=True))


def verify_execution_recovery():
    """核对真实 API 进程重启把遗留执行收敛为 unconfirmed。"""
    evidence = json.loads((REPORT / "browser.json").read_text())
    account = json.loads((REPORT / "account.json").read_text())
    interrupted = evidence["restart_inputs"]["interrupted_execution"]
    engine = create_control_engine(
        ControlDatabaseConfig.from_environment(require_migrator=False)
    )
    try:
        with engine.connect() as connection:
            versions = set(
                connection.execute(
                    text("SELECT version FROM schema_migrations")
                ).scalars()
            )
            required_versions = {
                f"chatbi-control-v{version}" for version in range(1, 6)
            }
            if not required_versions.issubset(versions):
                raise RuntimeError("重启后Control DB版本标记不完整")
            row = (
                connection.execute(
                    text(
                        """SELECT history.last_success_turn_id::text AS last_success_turn_id,
                    history.active_turn_id::text AS active_turn_id, turn.status AS turn_status,
                    turn.snapshot IS NULL AS snapshot_absent, execution.status AS execution_status,
                    execution.id::text AS execution_id
                    FROM history_records AS history
                    JOIN history_turns AS turn ON turn.history_id=history.id AND turn.id=:turn
                    JOIN history_executions AS execution ON execution.history_id=history.id
                        AND execution.turn_id=turn.id
                    WHERE history.id=:history AND history.owner_user_id=:owner"""
                    ),
                    {
                        "history": interrupted["history_id"],
                        "turn": interrupted["turn_id"],
                        "owner": account["id"],
                    },
                )
                .mappings()
                .one_or_none()
            )
        if row is None or row["execution_id"] != interrupted["execution_id"]:
            raise RuntimeError("重启后的执行记录身份不匹配")
        if (
            row["execution_status"] != "unconfirmed"
            or row["turn_status"] != "unconfirmed"
            or row["active_turn_id"] is not None
            or not row["snapshot_absent"]
            or row["last_success_turn_id"] != interrupted["prior_success_turn_id"]
        ):
            raise RuntimeError("重启未保留结果未确认与上一成功轮次边界")
        result = {
            "execution_id": row["execution_id"],
            "execution_status": row["execution_status"],
            "turn_status": row["turn_status"],
            "active_turn_cleared": row["active_turn_id"] is None,
            "unconfirmed_snapshot_absent": row["snapshot_absent"],
            "prior_success_turn_preserved": row["last_success_turn_id"]
            == interrupted["prior_success_turn_id"],
            "schema_versions": sorted(versions),
        }
        (REPORT / "recovery.json").write_text(json.dumps(result))
        evidence["execution_recovery"] = result
        (REPORT / "browser.json").write_text(
            json.dumps(evidence, ensure_ascii=False, indent=2)
        )
    finally:
        engine.dispose()
    print("真实API重启后的未确认执行、上一成功轮次与Schema版本检查通过。")


def _compact(value):
    return "".join(unicodedata.normalize("NFKC", str(value)).split())


def _verify_png(path: Path):
    raw = path.read_bytes()
    if not raw.startswith(b"\x89PNG\r\n\x1a\n"):
        raise RuntimeError("PNG 下载文件签名无效")
    offset = 8
    compressed = bytearray()
    width = height = 0
    ended = False
    while offset < len(raw):
        if offset + 12 > len(raw):
            raise RuntimeError("PNG chunk 被截断")
        length = struct.unpack(">I", raw[offset : offset + 4])[0]
        kind = raw[offset + 4 : offset + 8]
        end = offset + 12 + length
        if end > len(raw):
            raise RuntimeError("PNG chunk 长度超出文件")
        data = raw[offset + 8 : offset + 8 + length]
        expected_crc = struct.unpack(">I", raw[offset + 8 + length : end])[0]
        if zlib.crc32(kind + data) & 0xFFFFFFFF != expected_crc:
            raise RuntimeError("PNG chunk CRC 校验失败")
        if kind == b"IHDR":
            width, height, depth, color, compression, filtering, interlace = (
                struct.unpack(">IIBBBBB", data)
            )
            if (
                not width
                or not height
                or depth != 8
                or color not in {2, 6}
                or compression
                or filtering
                or interlace
            ):
                raise RuntimeError("PNG 尺寸或像素编码不支持")
        elif kind == b"IDAT":
            compressed.extend(data)
        elif kind == b"IEND":
            ended = True
            if end != len(raw):
                raise RuntimeError("PNG IEND 后仍有数据")
            break
        offset = end
    if not ended or not width or not height:
        raise RuntimeError("PNG 文件缺少完整结束标记")
    decoded = zlib.decompress(compressed)
    channels = {2: 3, 6: 4}[color]
    if len(decoded) != height * (1 + width * channels):
        raise RuntimeError("PNG 像素流不完整")
    return {"width": width, "height": height, "decoded_bytes": len(decoded)}


def _same_xlsx_value(actual, expected):
    if expected is None:
        return actual == "〈NULL〉"
    if expected == "":
        return actual == "〈空字符串〉"
    if isinstance(expected, bool):
        return actual is expected
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        try:
            return Decimal(str(actual)) == Decimal(str(expected))
        except (InvalidOperation, TypeError, ValueError):
            return str(actual) == str(expected)
    return actual == expected


def _verify_xlsx(path: Path, capture: dict):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != ["原始数据", "结果说明"]:
            raise RuntimeError("XLSX 工作表不符合导出 Contract")
        data = workbook["原始数据"]
        formulas = [
            cell.coordinate
            for row in data.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        if formulas:
            raise RuntimeError("XLSX 数据表包含公式")
        expected = capture.get("expected_snapshot")
        if expected:
            columns = expected["columns"]
            if [
                data.cell(1, index + 1).value for index in range(len(columns))
            ] != columns:
                raise RuntimeError("XLSX 列标题与已保存快照不一致")
            for row_index, row in enumerate(expected["rows"], start=2):
                actual = [
                    data.cell(row_index, index + 1).value
                    for index in range(len(columns))
                ]
                if len(row) != len(actual) or any(
                    not _same_xlsx_value(value, source)
                    for value, source in zip(actual, row, strict=True)
                ):
                    raise RuntimeError("XLSX 行数据与已保存快照不一致")
        return {
            "sheets": workbook.sheetnames,
            "rows": data.max_row - 1,
            "columns": data.max_column,
        }
    finally:
        workbook.close()


def _verify_pdf(path: Path, expected: dict):
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted or len(reader.pages) < 2:
        raise RuntimeError("PDF 无法打开或没有多页正文")
    text = _compact("\n".join(page.extract_text() or "" for page in reader.pages))
    required = [expected["question"]]
    report = expected["report"]
    required.extend(
        [report["title"], report["executive_summary"], report["trend_judgment"]]
    )
    required.extend(
        report["key_findings"] + report["root_causes"] + report["action_suggestions"]
    )
    attribution = report.get("attribution") or {}
    for key in (
        "metric_name",
        "comparison_period",
        "current_period",
        "comparison_value",
        "current_value",
        "total_change",
    ):
        if key in attribution:
            required.append(attribution[key])
    for product in attribution.get("products", []):
        required.extend((product["product_name"], product["change"]))
        required.extend(factor["name"] for factor in product["factors"])
        required.extend(factor["amount"] for factor in product["factors"])
    for task in expected["task_results"]:
        required.append(task["task_id"])
        required.extend(task["columns"])
        metadata = task.get("result_metadata") or {}
        for column in metadata.get("columns", []):
            required.extend(
                column[key]
                for key in ("semantic_name", "name", "definition")
                if isinstance(column.get(key), str) and column[key]
            )
            unit = column.get("unit")
            if isinstance(unit, dict) and isinstance(unit.get("label"), str):
                required.append(unit["label"])
        scope = metadata.get("scope") or {}
        required.extend(
            value
            for value in (scope.get("time") or {}).values()
            if isinstance(value, str) and value
        )
        for row in task["rows"]:
            required.extend(
                "NULL（无数据）"
                if cell is None
                else "空字符串"
                if cell == ""
                else "true"
                if cell is True
                else "false"
                if cell is False
                else str(cell)
                for cell in row
            )
    missing = [
        value for value in required if _compact(value) and _compact(value) not in text
    ]
    if missing:
        raise RuntimeError(f"PDF 可选取文本缺少快照字段（{len(missing)} 项）")
    for private in (expected["request_id"], expected["analysis_run_id"]):
        if _compact(private) in text:
            raise RuntimeError("PDF 包含私有运行标识")
    return {
        "pages": len(reader.pages),
        "checked_values": len(required),
        "text_characters": len(text),
    }


def verify_exports():
    evidence = json.loads((REPORT / "browser.json").read_text())
    captures = evidence.get("exports")
    if not isinstance(captures, list) or not captures:
        raise RuntimeError("浏览器报告没有导出证据")
    expected_pdf = evidence.get("analysis_pdf_expected")
    results = []
    for capture in captures:
        path = REPORT / capture["file"]
        if not path.is_file():
            raise RuntimeError("浏览器下载文件缺失")
        raw = path.read_bytes()
        if (
            len(raw) != capture["size"]
            or hashlib.sha256(raw).hexdigest() != capture["sha256"]
        ):
            raise RuntimeError("浏览器下载文件 hash / 大小与记录不一致")
        if capture["format"] == "xlsx":
            parsed = _verify_xlsx(path, capture)
        elif capture["format"] == "png":
            parsed = _verify_png(path)
        elif capture["format"] == "pdf":
            if not expected_pdf:
                raise RuntimeError("PDF 缺少真实分析快照对照值")
            parsed = _verify_pdf(path, expected_pdf)
        else:
            raise RuntimeError("浏览器报告含有未知导出格式")
        results.append(
            {
                "file": path.name,
                "format": capture["format"],
                "size": len(raw),
                "sha256": hashlib.sha256(raw).hexdigest(),
                "source_kind": capture["source"]["kind"],
                "parsed": parsed,
            }
        )
    formats = {item["format"] for item in results}
    if not {"xlsx", "png", "pdf"}.issubset(formats):
        raise RuntimeError("真实浏览器验收没有覆盖三种文件格式")
    sources = {item["source_kind"] for item in results}
    if not {"history_turn", "saved_result"}.issubset(sources):
        raise RuntimeError("真实浏览器验收没有覆盖历史与独立成果来源")
    by_format = {
        file_format: {
            item["source_kind"] for item in results if item["format"] == file_format
        }
        for file_format in ("xlsx", "png", "pdf")
    }
    if any(
        not {"history_turn", "saved_result"}.issubset(kinds)
        for kinds in by_format.values()
    ):
        raise RuntimeError("真实浏览器验收没有逐格式覆盖历史与独立成果来源")
    after_restart = {
        item["format"] for item in results if item["file"].startswith("r5-restarted-")
    }
    if after_restart != {"xlsx", "png", "pdf"}:
        raise RuntimeError("服务重启后没有重新下载并覆盖三种格式")
    result = {"status": "passed", "files": results, "formats": sorted(formats)}
    (REPORT / "export-verification.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2)
    )
    print("隔离 Compose 的 XLSX / PNG / PDF 浏览器下载已由独立解析器核验。")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "action",
        choices=[
            "prepare",
            "cleanup",
            "admin",
            "persistence",
            "expire-analysis",
            "execution-recovery",
            "verify-exports",
            "acceptance-snapshot",
        ],
    )
    args = parser.parse_args()
    if args.action == "prepare":
        prepare()
    elif args.action == "cleanup":
        cleanup()
    elif args.action == "admin":
        initialize_admin()
    elif args.action == "expire-analysis":
        expire_analysis()
    elif args.action == "execution-recovery":
        verify_execution_recovery()
    elif args.action == "verify-exports":
        verify_exports()
    elif args.action == "acceptance-snapshot":
        acceptance_snapshot()
    else:
        verify_persistence()
